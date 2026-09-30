"""Encrypted Excel storage for ONE user: one file per month (Sana | Vaqt | Summa | Izoh).

Every user gets their own SpendingStorage pointing at their own directory,
so one user's data is never read or written through another user's instance.
Files are .xlsx workbooks encrypted with the user's data key (see vault.py);
nothing readable is ever written to disk.

Callers work with a single `datetime`; it is split into the date and time
columns on write and combined again on read.
"""
from __future__ import annotations

import io
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Iterator

from cryptography.fernet import Fernet

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

HEADERS = ("Sana", "Vaqt", "Summa", "Izoh")
WIDTHS = (12, 8, 16, 80)
DAY_HEADERS = ("Sana", "Jami")
DAY_WIDTHS = (12, 16)
DATE_FORMAT = "DD.MM.YYYY"
TIME_FORMAT = "HH:MM"
TOP = Alignment(vertical="top")
WRAP_TOP = Alignment(vertical="top", wrap_text=True)

Spending = tuple[datetime, float, str]


def amount_format(amount: float) -> str:
    return "#,##0" if float(amount).is_integer() else "#,##0.00"


def iter_months(start: date, end: date) -> Iterator[tuple[int, int]]:
    year, month = start.year, start.month
    while (year, month) <= (end.year, end.month):
        yield year, month
        year, month = (year + 1, 1) if month == 12 else (year, month + 1)


def _setup_sheet(ws: Worksheet, title: str, headers: tuple[str, ...], widths: tuple[int, ...]) -> None:
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
    for col, width in zip("ABCDEFG", widths):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"


def _read_when(day_value, time_value) -> datetime | None:
    """Combine date and time cell values back into one datetime."""
    if isinstance(day_value, datetime):
        day_value = day_value.date()
    if not isinstance(day_value, date):
        return None
    if isinstance(time_value, datetime):
        time_value = time_value.time()
    if not isinstance(time_value, time):
        time_value = time()
    return datetime.combine(day_value, time_value)


def _write_amount(ws: Worksheet, row: int, amount: float) -> None:
    cell = ws.cell(row=row, column=3, value=amount)
    cell.number_format = amount_format(amount)
    cell.alignment = TOP


def _write_reason(ws: Worksheet, row: int, reason: str) -> None:
    cell = ws.cell(row=row, column=4, value=reason)
    # Force plain text so reasons starting with "=" are never treated as formulas.
    cell.data_type = "s"
    cell.alignment = WRAP_TOP


def _write_row(ws: Worksheet, row: int, when: datetime, amount: float, reason: str) -> None:
    date_cell = ws.cell(row=row, column=1, value=when.date())
    date_cell.number_format = DATE_FORMAT
    date_cell.alignment = TOP
    time_cell = ws.cell(row=row, column=2, value=when.time().replace(second=0, microsecond=0))
    time_cell.number_format = TIME_FORMAT
    time_cell.alignment = TOP
    _write_amount(ws, row, amount)
    _write_reason(ws, row, reason)


def _read_row(ws: Worksheet, row: int) -> Spending:
    return (
        _read_when(ws.cell(row=row, column=1).value, ws.cell(row=row, column=2).value),
        ws.cell(row=row, column=3).value or 0,
        ws.cell(row=row, column=4).value or "",
    )


SUFFIX = ".enc"


def shift_month(year: int, month: int, delta: int) -> tuple[int, int]:
    index = year * 12 + (month - 1) + delta
    return index // 12, index % 12 + 1


def purge_old(user_dir: Path, keep_from: tuple[int, int]) -> int:
    """Delete month files older than `keep_from` (year, month). Needs no key.

    Returns the number of deleted files.
    """
    deleted = 0
    if not user_dir.is_dir():
        return 0
    for path in user_dir.iterdir():
        if path.suffix not in (SUFFIX, ".xlsx"):
            continue
        try:
            year, month = map(int, path.stem.split("-"))
        except ValueError:
            continue
        if (year, month) < keep_from:
            path.unlink()
            deleted += 1
    return deleted


def encrypt_legacy(user_dir: Path, cipher: Fernet) -> int:
    """Encrypt any plaintext .xlsx month files left by older versions, then remove them."""
    converted = 0
    for path in list(user_dir.glob("*.xlsx")):
        target = path.with_suffix(SUFFIX)
        if not target.exists():
            target.write_bytes(cipher.encrypt(path.read_bytes()))
        path.unlink()
        converted += 1
    return converted


class SpendingStorage:
    def __init__(self, user_dir: Path, cipher: Fernet):
        self.dir = user_dir
        self._cipher = cipher

    def _load(self, path: Path, read_only: bool = False) -> Workbook:
        data = self._cipher.decrypt(path.read_bytes())
        return load_workbook(io.BytesIO(data), read_only=read_only)

    def _save(self, wb: Workbook, path: Path) -> None:
        buf = io.BytesIO()
        wb.save(buf)
        # Write to a temp file then rename, so a crash never leaves a half-written file.
        tmp = path.with_name(path.name + ".tmp")
        tmp.write_bytes(self._cipher.encrypt(buf.getvalue()))
        tmp.replace(path)

    def path_for(self, year: int, month: int) -> Path:
        return self.dir / f"{year:04d}-{month:02d}{SUFFIX}"

    def exists(self, year: int, month: int) -> bool:
        return self.path_for(year, month).exists()

    def ensure(self, year: int, month: int) -> Path:
        path = self.path_for(year, month)
        if not path.exists():
            self.dir.mkdir(parents=True, exist_ok=True)
            wb = Workbook()
            _setup_sheet(wb.active, f"{month:02d}.{year}", HEADERS, WIDTHS)
            self._save(wb, path)
        return path

    def entries(self, year: int, month: int) -> list[Spending]:
        path = self.path_for(year, month)
        if not path.exists():
            return []
        wb = self._load(path, read_only=True)
        try:
            return [
                (_read_when(d, t), amount or 0, reason or "")
                for d, t, amount, reason in wb.active.iter_rows(
                    min_row=2, max_col=4, values_only=True
                )
                if d is not None or amount is not None
            ]
        finally:
            wb.close()

    def entries_between(self, start: date, end: date) -> list[Spending]:
        """All spendings with start <= date <= end, in file order."""
        return [
            entry
            for year, month in iter_months(start, end)
            for entry in self.entries(year, month)
            if entry[0] is not None and start <= entry[0].date() <= end
        ]

    def append(self, when: datetime, amount: float, reason: str) -> int:
        """Add a spending at the end of its month's file. Returns its number (1-based)."""
        path = self.ensure(when.year, when.month)
        wb = self._load(path)
        ws = wb.active
        row = ws.max_row + 1
        _write_row(ws, row, when, amount, reason)
        self._save(wb, path)
        return row - 1

    def insert_sorted(self, when: datetime, amount: float, reason: str) -> int:
        """Add a spending in chronological position. Returns its number (1-based)."""
        path = self.ensure(when.year, when.month)
        wb = self._load(path)
        ws = wb.active
        row = ws.max_row + 1
        for r in range(2, ws.max_row + 1):
            existing = _read_row(ws, r)[0]
            if existing is not None and existing > when:
                row = r
                ws.insert_rows(r)
                break
        _write_row(ws, row, when, amount, reason)
        self._save(wb, path)
        return row - 1

    def update(self, year: int, month: int, number: int, amount: float, reason: str) -> Spending | None:
        """Replace amount and reason of spending `number`. Returns the old one, or None."""
        path = self.path_for(year, month)
        if not path.exists():
            return None
        wb = self._load(path)
        ws = wb.active
        row = number + 1
        if number < 1 or row > ws.max_row:
            return None
        old = _read_row(ws, row)
        _write_amount(ws, row, amount)
        _write_reason(ws, row, reason)
        self._save(wb, path)
        return old

    def get(self, year: int, month: int, number: int) -> Spending | None:
        entries = self.entries(year, month)
        return entries[number - 1] if 1 <= number <= len(entries) else None

    def delete(self, year: int, month: int, number: int) -> Spending | None:
        """Delete spending `number`. Returns it, or None if missing."""
        path = self.path_for(year, month)
        if not path.exists():
            return None
        wb = self._load(path)
        ws = wb.active
        row = number + 1
        if number < 1 or row > ws.max_row:
            return None
        removed = _read_row(ws, row)
        ws.delete_rows(row)
        self._save(wb, path)
        return removed

    def export_range(self, start: date, end: date, title: str) -> tuple[bytes, int, float]:
        """A workbook with all spendings in [start, end], a JAMI row and per-day totals.

        Returns (xlsx bytes, count, total).
        """
        rows = self.entries_between(start, end)
        wb = Workbook()
        ws = wb.active
        _setup_sheet(ws, title, HEADERS, WIDTHS)
        for when, amount, reason in rows:
            _write_row(ws, ws.max_row + 1, when, amount, reason)
        total = sum(amount for _, amount, _ in rows)
        label = ws.cell(row=ws.max_row + 2, column=1, value="JAMI")
        label.font = Font(bold=True)
        total_cell = ws.cell(row=label.row, column=3, value=total)
        total_cell.font = Font(bold=True)
        total_cell.number_format = amount_format(total)

        per_day: dict[date, float] = {}
        for when, amount, _ in rows:
            per_day[when.date()] = per_day.get(when.date(), 0) + amount
        days = wb.create_sheet()
        _setup_sheet(days, "Kunlik jami", DAY_HEADERS, DAY_WIDTHS)
        day = start
        while day <= end:
            if day in per_day:
                days.append((day, per_day[day]))
                days.cell(row=days.max_row, column=1).number_format = DATE_FORMAT
                days.cell(row=days.max_row, column=2).number_format = amount_format(per_day[day])
            day += timedelta(days=1)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue(), len(rows), total
